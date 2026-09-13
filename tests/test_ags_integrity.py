"""The fixture corpus is identical to the frontend merge contract tests."""
import io
import json
from pathlib import Path
import tempfile
import unittest

from boreholeai._merge._ags import AgsMergeConflict, merge_ags_files
from boreholeai._merge import merge_results
from boreholeai._merge._excel import merge_excel_files
from boreholeai._merge._json import merge_json_files
from openpyxl import Workbook, load_workbook


class AgsIntegrityTests(unittest.TestCase):
    def test_shared_merge_fixtures(self):
        cases = json.loads((Path(__file__).parent / 'fixtures/ags-merge-integrity.json').read_text())
        for case in cases:
            with self.subTest(case=case['name']), tempfile.TemporaryDirectory() as folder:
                paths = []
                for index, content in enumerate(case['inputs']):
                    path = Path(folder) / f'{index}.ags'
                    path.write_text(content)
                    paths.append(path)
                if 'error' in case:
                    with self.assertRaisesRegex(AgsMergeConflict, case['error']):
                        merge_ags_files(paths)
                else:
                    self.assertEqual(merge_ags_files(paths).replace('\r\n', '\n'), case['expected'])

    def test_conflict_retains_originals_and_removes_stale_merge(self):
        cases = json.loads((Path(__file__).parent / 'fixtures/ags-merge-integrity.json').read_text())
        case = next(c for c in cases if c['name'] == 'unit_conflict')
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            dirs = []
            for i, content in enumerate(case['inputs']):
                directory = root / str(i)
                directory.mkdir()
                (directory / 'Borehole_ags4.ags').write_text(content)
                dirs.append(directory)
            out = root / 'merged'
            out.mkdir()
            (out / 'Borehole_ags4_merged.ags').write_text('STALE')
            result = merge_results(dirs, out)
            self.assertFalse((out / 'Borehole_ags4_merged.ags').exists())
            self.assertTrue(any('UNIT conflict' in warning for warning in result.warnings))
            for i, content in enumerate(case['inputs'], 1):
                self.assertEqual((out / 'original_ags' / f'{i}_Borehole_ags4.ags').read_text(), content)

    def test_json_excel_preserve_qualified_results_and_metadata(self):
        rows = [
            {"Hole_ID": "BH1", "from": 2, "Cu": value, "Cu_unit": unit,
             "source_text": str(value), "sample_ref": "001",
             "source_observation_id": f"observation-{index}"}
            for index, (value, unit) in enumerate([(2, "tsf"), (">200", "kPa"), ("<=50", None), ("290–340", "kPa")])
        ]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            json_paths, excel_paths = [], []
            for index, row in enumerate(rows):
                jp = root / f"{index}.json"
                jp.write_text(json.dumps({"test_data": {"cu": [row]}}))
                json_paths.append(jp)
                book = Workbook()
                sheet = book.active
                sheet.title = "Cu"
                # Reverse incoming heading order to exercise realignment.
                headings = list(row) if index % 2 else list(reversed(row))
                sheet.append(headings)
                sheet.append([row[h] for h in headings])
                xp = root / f"{index}.xlsx"
                book.save(xp)
                excel_paths.append(xp)
            self.assertEqual(json.loads(merge_json_files(json_paths))["test_data"]["cu"], rows)
            book = load_workbook(io.BytesIO(merge_excel_files(excel_paths)))
            values = list(book["Cu"].values)
            actual = [dict(zip(values[0], row)) for row in values[1:]]
            for expected, actual_row in zip(rows, actual):
                for key, value in expected.items():
                    self.assertEqual(actual_row[key], value)


if __name__ == '__main__':
    unittest.main()
